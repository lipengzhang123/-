import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "精益改善案例平台项目计划"

# 表头
headers = ["周次", "时间范围", "核心事项", "负责人", "备注"]
ws.append(headers)

# 数据
data = [
    ("W37", "09.07-09.13", 
     "①梳理近两年全部黑/绿带项目，确定项目改善类别\n②定稿案例分类、标签、一页纸模板及素材归档规范\n③输出AI平台功能、三级鉴权、AI能力完整需求清单", "", ""),
    ("W38", "09.14-09.20", 
     "①固化所有案例标准、平台需求，完成需求终审锁版\n②技术侧启动平台前端、后端、鉴权模块开发筹备工作", "", ""),
    ("W39", "09.21-09.27", 
     "①对接各项目负责人，批量补齐案例缺失数据、素材、实操及避坑内容\n②推进平台前端页面、基础鉴权功能开发", "", "4天工作日"),
    ("W40", "09.28-10.04", 
     "①完成所有带级案例整编、校对、分类打标签及网盘归档，补齐五大模块外部标杆案例\n②完成平台AI功能、权限隔离、检索下载等主体功能开发", "", "3天工作日（含国庆）"),
    ("W41", "10.05-10.11", 
     "①完成平台服务部署、环境配置、内网访问开通\n②全量案例数据批量导入平台，完成基础数据校验\n③开展全功能首轮联调与BUG初修", "", "3天工作日"),
    ("W42", "10.12-10.18", 
     "①完成平台所有功能、权限、自动化PDF链路调试优化\n②完成页面样式、使用体验统一整改，平台达到可用标准", "", ""),
    ("W43", "10.19-10.25", 
     "①组织各分子公司精益负责人开展平台内测，收集问题清单\n②闭环修复所有功能、数据、权限问题，优化平台体验", "", ""),
    ("W44", "10.26-11.01", 
     "①输出平台操作、案例提报、运维规范文档\n②完成全项目最终验收，案例库、AI平台正式上线启用", "", "")
]

for row in data:
    ws.append(row)

# 设置列宽
ws.column_dimensions['A'].width = 8
ws.column_dimensions['B'].width = 14
ws.column_dimensions['C'].width = 65
ws.column_dimensions['D'].width = 12
ws.column_dimensions['E'].width = 14

# 表头样式
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="0052CC", end_color="0052CC", fill_type="solid")
header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

for col in range(1, len(headers)+1):
    cell = ws.cell(row=1, column=col)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_alignment
    cell.border = thin_border

# 数据行样式
data_alignment = Alignment(vertical="top", wrap_text=True)
alt_fill = PatternFill(start_color="E6F0FF", end_color="E6F0FF", fill_type="solid")

for row_idx in range(2, len(data)+2):
    for col in range(1, len(headers)+1):
        cell = ws.cell(row=row_idx, column=col)
        cell.alignment = data_alignment
        cell.border = thin_border
        if row_idx % 2 == 0:
            cell.fill = alt_fill

# 保存
output_path = r"C:\Users\22104\Desktop\精益改善案例平台项目计划.xlsx"
wb.save(output_path)
print(f"已生成: {output_path}")
